import os
import time
import fitz 
from pptx import Presentation 
import google.generativeai as genai
from pinecone import Pinecone
from dotenv import load_dotenv
import re

# --- CONFIGURATION ---
load_dotenv()
genai.configure(api_key=os.getenv("AIzaSyDCUHFPcdvSL9kjAvWIqAhPV8o-GZr3YF4"))
PINECONE_API_KEY = os.getenv("pcsk_QoigX_GrE3DY7B9Jcj4ZDeinkhDpniqiWZEVx1GK4TogXDxMQZq5q2iwN5NwTjWndNF3L")
PINECONE_INDEX_NAME = 'sih-rag-index'
DOCS_FOLDER = './source_documents'

print("Initializing Pinecone connection...")
pc = Pinecone(api_key=PINECONE_API_KEY)

# --- 1. CREATE PINECONE INDEX (if it doesn't exist) ---
if PINECONE_INDEX_NAME not in pc.list_indexes().names():
    print(f"Creating new Pinecone index: {PINECONE_INDEX_NAME}")
    pc.create_index(
        name=PINECONE_INDEX_NAME,
        dimension=768, # Dimension for the Google embedding model 'text-embedding-004'
        metric='cosine'
    )
index = pc.Index(PINECONE_INDEX_NAME)
print("Pinecone index is ready.")


def get_text_from_doc(file_path):
    """Extracts text from PDF or PPTX files."""
    if file_path.lower().endswith(".pdf"):
        try:
            doc = fitz.open(file_path)
            text = "".join(page.get_text() for page in doc)
            return text
        except Exception as e:
            print(f"   [!] Error reading PDF {os.path.basename(file_path)}: {e}")
            return ""
    elif file_path.lower().endswith(".pptx"):
        try:
            pres = Presentation(file_path)
            text = "\n".join(shape.text for slide in pres.slides for shape in slide.shapes if hasattr(shape, "text"))
            return text
        except Exception as e:
            print(f"   [!] Error reading PPTX {os.path.basename(file_path)}: {e}")
            return ""
    return ""

def chunk_text(text):
    """Splits text into meaningful chunks based on paragraphs."""
    text = re.sub(r'\n\s*\n', '\n\n', text)
    chunks = text.split('\n\n')
    return [chunk.strip() for chunk in chunks if len(chunk.strip()) > 150]

def create_embeddings_and_index():
    """Processes all documents, creates embeddings, and upserts to Pinecone."""
    print(f"\nStarting document search in '{os.path.abspath(DOCS_FOLDER)}'...")
    total_chunks_indexed = 0
    batch_size = 50 

    documents_to_index = []
    # --- UPGRADED SCRIPT TO SEARCH ALL SUB-FOLDERS ---
    # os.walk() goes through every folder and file recursively
    for root, dirs, files in os.walk(DOCS_FOLDER):
        print(f"\nSearching in folder: {root}")
        for filename in files:
            if filename.lower().endswith((".pdf", ".pptx")):
                print(f"  -> Found document: {filename}")
                file_path = os.path.join(root, filename)
                text = get_text_from_doc(file_path)
                if not text:
                    print(f"     - Skipping document, no text found.")
                    continue
                
                chunks = chunk_text(text)
                print(f"     - Split into {len(chunks)} chunks.")
                for i, chunk in enumerate(chunks):
                    documents_to_index.append({
                        'id': f"{filename}-{i}",
                        'content': chunk,
                        'metadata': {'source': filename, 'text': chunk}
                    })
            else:
                print(f"  -> Skipping non-document file: {filename}")


    if not documents_to_index:
        print("\n[!] No documents found or processed. Please check the 'source_documents' folder and file types.")
        return

    print(f"\n--- Total of {len(documents_to_index)} chunks found. Now creating embeddings in batches. ---")

    # Process in batches
    for i in range(0, len(documents_to_index), batch_size):
        batch = documents_to_index[i:i+batch_size]
        contents = [doc['content'] for doc in batch]
        
        print(f"\n--- Processing Batch {i//batch_size + 1} of { (len(documents_to_index) + batch_size - 1) // batch_size } ---")
        try:
            print(f"   Creating embeddings for {len(contents)} chunks with Google AI...")
            result = genai.embed_content(
                model="models/text-embedding-004",
                content=contents,
                task_type="RETRIEVAL_DOCUMENT"
            )
            embeddings = result["embedding"]
            
            vectors_to_upsert = []
            for j, doc in enumerate(batch):
                vectors_to_upsert.append((doc['id'], embeddings[j], doc['metadata']))
            
            print(f"   Uploading {len(vectors_to_upsert)} vectors to Pinecone...")
            index.upsert(vectors=vectors_to_upsert)
            total_chunks_indexed += len(vectors_to_upsert)

        except Exception as e:
            print(f"   [!] ERROR processing batch: {e}")
        
        if len(documents_to_index) > i + batch_size:
            print("   Batch complete. Waiting for 5 seconds to avoid rate limits...")
            time.sleep(5)

    print(f"\n✅ Indexing complete! A total of {total_chunks_indexed} text chunks have been indexed in Pinecone.")

if __name__ == '__main__':
    create_embeddings_and_index()

    

