import os
import io
import json
import traceback
from flask import Flask, request, jsonify, render_template
from google.cloud import dialogflow_v2 as dialogflow
from google.cloud import firestore
from pinecone import Pinecone
import google.generativeai as genai
from dotenv import load_dotenv
from langdetect import detect
from google.oauth2.credentials import Credentials

# --- IMPORTS for Drive, Vision, STT, TTS, and Document processing ---
from google.cloud import texttospeech
from google.cloud import speech
from google.cloud import vision
import base64
import fitz  # PyMuPDF
from pptx import Presentation
import docx
from openpyxl import load_workbook
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload
# --------------------------------------------------------------------

# --- 1. INITIALIZATION AND CONFIGURATION ---
load_dotenv()
app = Flask(__name__, static_folder='static', template_folder='templates')

try:
    # --- Google Cloud Services Authentication for Vercel ---
    firebase_creds_json = os.environ.get('FIREBASE_CREDENTIALS')
    
    if firebase_creds_json:
        # Vercel environment: Write the secret string to a temporary serverless file
        tmp_key_path = '/tmp/serviceAccountKey.json'
        with open(tmp_key_path, 'w') as f:
            f.write(firebase_creds_json)
        os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = tmp_key_path
    else:
        # Local development environment: Use your physical file
        os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = 'serviceAccountKey.json'

    GCLOUD_PROJECT = os.getenv("GCLOUD_PROJECT")
    db = firestore.Client()
    
    # --- Pinecone ---
    PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
    PINECONE_INDEX_NAME = 'sih-rag-index'
    pc = Pinecone(api_key=PINECONE_API_KEY)
    pinecone_index = pc.Index(PINECONE_INDEX_NAME)

    # --- Gemini ---
    GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
    genai.configure(api_key=GEMINI_API_KEY)
    gemini_generative_model = genai.GenerativeModel('gemini-1.5-flash')
    gemini_embedding_model = "models/text-embedding-004"

    # --- Initialize all Google Cloud Clients ---
    speech_client = speech.SpeechClient()
    tts_client = texttospeech.TextToSpeechClient()
    vision_client = vision.ImageAnnotatorClient() 
    
    print("✅ All services (Dialogflow, Firestore, Pinecone, Gemini, Speech, TTS, Vision) configured successfully.")

except Exception as e:
    print(f"❌ Error during configuration: {e}")
    traceback.print_exc()

# --- 2. WEB ROUTES ---
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/webhook', methods=['POST'])
def webhook():
    req = request.get_json(force=True)
    user_query = req.get('queryResult', {}).get('queryText', '')
    session_id = req.get('session', 'default-session').split('/')[-1]
    
    print(f"\n➡️ [V1] Received query: '{user_query}'")

    try:
        try:
            detected_language_code = detect(user_query)
        except: 
           detected_language_code = 'en'
        print(f"   ✅ Server auto-detected language: '{detected_language_code}'")

        dialogflow_response = detect_intent_texts(GCLOUD_PROJECT, session_id, user_query, detected_language_code)
        intent_name = dialogflow_response.query_result.intent.display_name
        print(f"   ✅ Dialogflow matched intent: '{intent_name}'")
        
        response_text = ""
        if intent_name != 'Default Fallback Intent':
            print("   ➡️ Routing to Admin Brain (Firestore)...")
            response_text = get_response_from_firestore(intent_name, detected_language_code)
        else:
            print("   🤔 No specific intent matched. Routing to Expert Brain (Gemini RAG)...")
            response_text = get_gemini_rag_response(user_query)
        
        return jsonify({'fulfillmentText': response_text})

    except Exception as e:
        print(f"❌ An error occurred in the webhook.")
        traceback.print_exc()
        return jsonify({'fulfillmentText': "Sorry, I encountered a server error."})

@app.route('/webhook_v2', methods=['POST'])
def webhook_v2():
    req = request.get_json(force=True)
    
    user_query = req.get('queryInput', {}).get('text', {}).get('text', '')
    session_id = req.get('session', 'default-session').split('/')[-1]
    language_code = req.get('queryInput', {}).get('text', {}).get('languageCode', 'en')
    
    file_context = None
    file_name = None
    if 'queryParams' in req and 'payload' in req['queryParams']:
        file_context = req['queryParams']['payload'].get('file_context')
        file_name = req['queryParams']['payload'].get('file_name')

    print(f"\n➡️ [V2] Received query: '{user_query}'")
    if file_name:
        print(f"   📎 With context from file: {file_name}")

    try:
        if file_context:
            print("   📄 File context found. Routing directly to Expert Brain (Gemini RAG)...")
            response_text = get_gemini_rag_response_with_context(user_query, file_context)
        else:
            dialogflow_response = detect_intent_texts(GCLOUD_PROJECT, session_id, user_query, language_code)
            intent_name = dialogflow_response.query_result.intent.display_name
            print(f"   ✅ Dialogflow matched intent: '{intent_name}'")
            
            if intent_name != 'Default Fallback Intent':
                print("   ➡️ Routing to Admin Brain (Firestore)...")
                response_text = get_response_from_firestore(intent_name, language_code)
            else:
                print("   🤔 No specific intent matched. Routing to Expert Brain (Gemini RAG)...")
                response_text = get_gemini_rag_response(user_query)
        
        return jsonify({'fulfillmentText': response_text})

    except Exception as e:
        print(f"❌ An error occurred in the webhook v2.")
        traceback.print_exc()
        return jsonify({'fulfillmentText': "Sorry, I encountered a server error."})


# --- VOICE ENDPOINTS ---
@app.route('/recognize', methods=['POST'])
def recognize():
    audio_data = request.json.get('audioData')
    language_code = request.json.get('languageCode', 'en-US')
    if not audio_data:
        return jsonify({'error': 'No audio data provided'}), 400
    try:
        audio_content = base64.b64decode(audio_data)
        audio = speech.RecognitionAudio(content=audio_content)
        config = speech.RecognitionConfig(
            encoding=speech.RecognitionConfig.AudioEncoding.WEBM_OPUS,
            sample_rate_hertz=48000,
            language_code=language_code,
            model="default"
        )
        print(f"\n🎤 Transcribing audio in language: {language_code}")
        response = speech_client.recognize(config=config, audio=audio)
        if response.results and response.results[0].alternatives:
            transcript = response.results[0].alternatives[0].transcript
            print(f"   ✅ Transcription successful: '{transcript}'")
            return jsonify({'transcript': transcript})
        else:
            print("   ⚠️ No transcription result.")
            return jsonify({'transcript': ''})
    except Exception as e:
        print(f"❌ Speech-to-Text Error: {e}")
        traceback.print_exc()
        return jsonify({'error': 'Failed to transcribe audio'}), 500

@app.route('/synthesize', methods=['POST'])
def synthesize():
    text = request.json.get('text')
    language_code = request.json.get('languageCode', 'en-US')
    if not text:
        return jsonify({'error': 'No text provided'}), 400
    try:
        synthesis_input = texttospeech.SynthesisInput(text=text)
        if language_code.startswith('hi'):
            voice_name = 'hi-IN-Wavenet-B'
        elif language_code.startswith('gu'):
            voice_name = 'gu-IN-Wavenet-A'
        else:
            voice_name = 'en-IN-Wavenet-A'
        voice = texttospeech.VoiceSelectionParams(language_code=language_code, name=voice_name)
        audio_config = texttospeech.AudioConfig(audio_encoding=texttospeech.AudioEncoding.MP3)
        print(f"\n🔊 Synthesizing speech for text: '{text[:30]}...' in language {language_code}")
        response = tts_client.synthesize_speech(input=synthesis_input, voice=voice, audio_config=audio_config)
        print("   ✅ Speech synthesis successful.")
        return jsonify({'audioContent': base64.b64encode(response.audio_content).decode('utf-8')})
    except Exception as e:
        print(f"❌ Text-to-Speech Error: {e}")
        traceback.print_exc()
        return jsonify({'error': 'Failed to synthesize speech'}), 500


# --- FILE PROCESSING ENDPOINTS ---
@app.route('/upload_file', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No selected file'}), 400
    
    try:
        content = file.read()
        mime_type = file.mimetype
        extracted_text = ""
        print(f"\n📄 Received local file: {file.filename} ({mime_type})")

        if 'image' in mime_type:
            extracted_text = extract_text_from_image(content)
        elif 'pdf' in mime_type:
            extracted_text = extract_text_from_pdf(content)
        elif 'presentationml' in mime_type: 
            extracted_text = extract_text_from_pptx(content)
        elif 'wordprocessingml' in mime_type: 
            extracted_text = extract_text_from_docx(content)
        elif 'spreadsheetml' in mime_type: 
            extracted_text = extract_text_from_xlsx(content)
        elif 'text' in mime_type:
            try:
                extracted_text = content.decode('utf-8')
            except UnicodeDecodeError:
                extracted_text = content.decode('latin-1')
        else:
            return jsonify({'error': f"Unsupported file type: {mime_type}. Please upload a supported document."}), 400
            
        print(f"   ✅ Extracted text successfully.")
        return jsonify({'extracted_text': extracted_text})
        
    except Exception as e:
        print(f"❌ File Upload Error: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/process_drive_file', methods=['POST'])
def process_drive_file():
    req_data = request.get_json()
    file_id = req_data.get('fileId')
    access_token = req_data.get('accessToken') 

    if not file_id or not access_token:
        return jsonify({'error': 'Missing fileId or accessToken'}), 400

    try:
        creds = Credentials(token=access_token)
        drive_service = build('drive', 'v3', credentials=creds)

        file_metadata = drive_service.files().get(fileId=file_id, fields='mimeType, name').execute()
        mime_type = file_metadata.get('mimeType')
        print(f"\n📄 Received Drive file from user: {file_metadata.get('name')} ({mime_type})")

        request_media = drive_service.files().get_media(fileId=file_id)
        file_content_io = io.BytesIO()
        downloader = MediaIoBaseDownload(file_content_io, request_media)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        
        content = file_content_io.getvalue()
        
        if 'image' in mime_type:
            extracted_text = extract_text_from_image(content)
        elif 'pdf' in mime_type:
            extracted_text = extract_text_from_pdf(content)
        elif 'presentationml' in mime_type:
            extracted_text = extract_text_from_pptx(content)
        elif 'wordprocessingml' in mime_type:
            extracted_text = extract_text_from_docx(content)
        elif 'spreadsheetml' in mime_type:
            extracted_text = extract_text_from_xlsx(content)
        elif 'text' in mime_type:
            extracted_text = content.decode('utf-8', errors='ignore')
        else:
            return jsonify({'error': f"Unsupported file type: {mime_type}."}), 400

        print(f"   ✅ Extracted text successfully.")
        return jsonify({'extracted_text': extracted_text})

    except Exception as e:
        print(f"❌ Drive Processing Error: {e}")
        traceback.print_exc()
        return jsonify({'error': 'Failed to process Google Drive file. The user token may be invalid or expired.'}), 500

# --- 3. HELPER FUNCTIONS ---
def extract_text_from_image(content):
    image = vision.Image(content=content)
    response = vision_client.text_detection(image=image)
    if response.error.message:
        raise Exception(f'Vision API Error: {response.error.message}')
    return response.full_text_annotation.text

def extract_text_from_pdf(content):
    text = ""
    with fitz.open(stream=content, filetype="pdf") as doc:
        for page in doc:
            text += page.get_text()
    return text

def extract_text_from_pptx(content):
    text = ""
    with io.BytesIO(content) as f:
        prs = Presentation(f)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    return text

def extract_text_from_docx(content):
    text = ""
    with io.BytesIO(content) as f:
        document = docx.Document(f)
        for para in document.paragraphs:
            text += para.text + "\n"
    return text

def extract_text_from_xlsx(content):
    text = ""
    with io.BytesIO(content) as f:
        workbook = load_workbook(f)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows():
                row_text = []
                for cell in row:
                    if cell.value is not None:
                        row_text.append(str(cell.value))
                if row_text:
                    text += " | ".join(row_text) + "\n"
    return text

def detect_intent_texts(project_id, session_id, text, language_code):
    session_client = dialogflow.SessionsClient()
    session = session_client.session_path(project_id, session_id)
    text_input = dialogflow.TextInput(text=text, language_code=language_code)
    query_input = dialogflow.QueryInput(text=text_input)
    response = session_client.detect_intent(request={"session": session, "query_input": query_input})
    return response

def get_response_from_firestore(intent_name, language_code):
    try:
        query = db.collection('faqs').where('intentName', '==', intent_name).limit(1)
        results = query.stream()
        for doc in results:
            doc_data = doc.to_dict()
            responses_map = doc_data.get('responseText', {})
            response = responses_map.get(language_code, responses_map.get('en'))
            if response:
                print(f"   ✅ Found response in Firestore for language '{language_code}'.")
                return response
            else:
                return "I have an answer for that, but not in your language yet."
        return f"I understand you're asking about '{intent_name}', but I don't have an answer for that yet in my database."
    except Exception as e:
        print(f"❌ Firestore Error: {e}")
        return "Sorry, I had trouble connecting to my database."

def get_gemini_rag_response(query):
    try:
        query_embedding = genai.embed_content(model=gemini_embedding_model, content=query, task_type="RETRIEVAL_QUERY")["embedding"]
        results = pinecone_index.query(vector=query_embedding, top_k=3, include_metadata=True)
        context = "\n".join([x['metadata']['text'] for x in results['matches']])
        augmented_prompt = f"Using ONLY this context from official documents:\n---\n{context}\n---\nAnswer the question: {query}"
        answer = gemini_generative_model.generate_content(augmented_prompt)
        return answer.text
    except Exception as e:
        print(f"❌ Gemini RAG Error: {e}")
        return "Sorry, I encountered an error with my advanced knowledge base. Please ensure billing is enabled for the project."

def get_gemini_rag_response_with_context(query, file_context):
    try:
        print("   🧠 Querying Pinecone for related context...")
        query_embedding = genai.embed_content(model=gemini_embedding_model, content=query, task_type="RETRIEVAL_QUERY")["embedding"]
        pinecone_results = pinecone_index.query(vector=query_embedding, top_k=3, include_metadata=True)
        pinecone_context = "\n".join([x['metadata']['text'] for x in pinecone_results['matches']])

        augmented_prompt = f"""You are an expert assistant. Answer the user's question based on the following context.

        First, prioritize the information from the "UPLOADED DOCUMENT CONTEXT", as it was provided directly by the user.
        Then, use the "ADDITIONAL CONTEXT FROM DATABASE" if it helps to provide a more complete answer.

        --- UPLOADED DOCUMENT CONTEXT ---
        {file_context}
        --- END OF UPLOADED DOCUMENT CONTEXT ---

        --- ADDITIONAL CONTEXT FROM DATABASE ---
        {pinecone_context}
        --- END OF ADDITIONAL CONTEXT FROM DATABASE ---

        User's Question: {query}
        """
        
        print("   ✨ Generating final answer with combined context...")
        answer = gemini_generative_model.generate_content(augmented_prompt)
        return answer.text
    except Exception as e:
        print(f"❌ Gemini RAG (with context) Error: {e}")
        traceback.print_exc()
        return "Sorry, I encountered an error with my advanced knowledge base while analyzing the document."